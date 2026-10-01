"""Turn files into text segments, each tagged with where it came from.

`extract()` returns a list of (location, text) pairs, e.g. ("page 3", "...").
The location is shown in search results so users can jump straight to the
right page / slide / sheet.
"""
from __future__ import annotations

import html
import io
import json
import os
import re
import shutil
import sys
from pathlib import Path

Segments = list[tuple[str, str]]

PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
PPTX_EXTS = {".pptx"}
XLSX_EXTS = {".xlsx", ".xlsm"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
HTML_EXTS = {".html", ".htm"}
NOTEBOOK_EXTS = {".ipynb"}
TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".yaml", ".yml",
    ".xml", ".tex", ".log", ".ini", ".cfg", ".toml",
    ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".c", ".h", ".cpp", ".hpp",
    ".cs", ".go", ".rs", ".rb", ".php", ".kt", ".swift", ".m", ".r", ".sql",
    ".sh", ".bat", ".ps1", ".css", ".scss",
}
SUPPORTED_EXTS = (
    PDF_EXTS | DOCX_EXTS | PPTX_EXTS | XLSX_EXTS | IMAGE_EXTS | HTML_EXTS
    | NOTEBOOK_EXTS | TEXT_EXTS
)

# A PDF page with less text than this is treated as scanned and OCR'd (if enabled).
_SCANNED_PAGE_CHARS = 25
_MAX_SHEET_ROWS = 20_000


# ---------------------------------------------------------------- OCR setup
_WINDOWS_TESSERACT = [
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
]


def configure_ocr(mode: str, tesseract_cmd: str = "") -> bool:
    """Point pytesseract at a tesseract binary. Returns True if OCR is usable."""
    if mode == "off":
        return False
    try:
        import pytesseract
    except ImportError:
        return False

    cmd = tesseract_cmd or shutil.which("tesseract") or ""
    if not cmd and sys.platform == "win32":
        cmd = next((p for p in _WINDOWS_TESSERACT if os.path.exists(p)), "")
    if cmd:
        pytesseract.pytesseract.tesseract_cmd = cmd
    try:
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        if mode == "on":
            raise RuntimeError(
                "DOCTRACE_OCR=on but tesseract was not found. Install it or set TESSERACT_CMD."
            ) from None
        return False


def _ocr_image(img) -> str:
    import pytesseract

    return pytesseract.image_to_string(img)


# ---------------------------------------------------------------- helpers
def _as_stream(src: str | bytes):
    return io.BytesIO(src) if isinstance(src, bytes) else src


def _read_bytes(src: str | bytes) -> bytes:
    if isinstance(src, bytes):
        return src
    with open(src, "rb") as fh:
        return fh.read()


def _decode(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


# ---------------------------------------------------------------- extractors
def _pdf(src, ocr: bool) -> Segments:
    import pymupdf

    pymupdf.TOOLS.mupdf_display_errors(False)
    doc = pymupdf.open(stream=src, filetype="pdf") if isinstance(src, bytes) else pymupdf.open(src)
    out = []
    with doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if ocr and len(text.strip()) < _SCANNED_PAGE_CHARS:
                from PIL import Image

                pix = page.get_pixmap(dpi=200)
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                text = _ocr_image(img)
            out.append((f"page {i}", text))
    return out


def _docx(src) -> Segments:
    import docx

    d = docx.Document(_as_stream(src))
    out: Segments = []
    heading, buf = "", []

    def flush():
        if buf:
            out.append((f"section \"{heading}\"" if heading else "", "\n".join(buf)))

    for p in d.paragraphs:
        style = (p.style.name if p.style is not None else "") or ""
        if style.lower().startswith(("heading", "title")) and p.text.strip():
            flush()
            heading, buf = p.text.strip()[:80], [p.text]
        elif p.text.strip():
            buf.append(p.text)
    flush()

    for t_i, table in enumerate(d.tables, start=1):
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        out.append((f"table {t_i}", "\n".join(rows)))
    return out


def _pptx(src) -> Segments:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def shape_text(shape) -> list[str]:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            return [t for s in shape.shapes for t in shape_text(s)]
        parts = []
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            parts.append(shape.text_frame.text)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return parts

    prs = Presentation(_as_stream(src))
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = [t for shape in slide.shapes for t in shape_text(shape)]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            parts.append(slide.notes_slide.notes_text_frame.text)
        out.append((f"slide {i}", "\n".join(p for p in parts if p.strip())))
    return out


def _xlsx(src) -> Segments:
    import openpyxl

    wb = openpyxl.load_workbook(_as_stream(src), read_only=True, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            rows = []
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= _MAX_SHEET_ROWS:
                    break
                cells = [str(v) for v in row if v is not None and str(v).strip()]
                if cells:
                    rows.append(" | ".join(cells))
            out.append((f"sheet \"{ws.title}\"", "\n".join(rows)))
    finally:
        wb.close()
    return out


_SCRIPT_STYLE = re.compile(r"<(script|style)\b.*?</\1>", re.S | re.I)
_TAG = re.compile(r"<[^>]+>")


def _html(src) -> Segments:
    raw = _SCRIPT_STYLE.sub(" ", _decode(_read_bytes(src)))
    raw = re.sub(r"<(br|/p|/div|/h\d|/li|/tr)\b[^>]*>", "\n", raw, flags=re.I)
    return [("", html.unescape(_TAG.sub(" ", raw)))]


def _notebook(src) -> Segments:
    nb = json.loads(_decode(_read_bytes(src)))
    out = []
    for i, cell in enumerate(nb.get("cells", []), start=1):
        source = cell.get("source", "")
        text = "".join(source) if isinstance(source, list) else str(source)
        out.append((f"cell {i}", text))
    return out


def _image(src, ocr: bool) -> Segments:
    if not ocr:
        return []
    from PIL import Image

    with Image.open(_as_stream(src)) as img:
        return [("", _ocr_image(img))]


def extract(src: str | bytes, ext: str, ocr: bool = False) -> Segments:
    """Extract (location, text) segments from a file path or raw bytes."""
    ext = ext.lower()
    if ext in PDF_EXTS:
        segs = _pdf(src, ocr)
    elif ext in DOCX_EXTS:
        segs = _docx(src)
    elif ext in PPTX_EXTS:
        segs = _pptx(src)
    elif ext in XLSX_EXTS:
        segs = _xlsx(src)
    elif ext in IMAGE_EXTS:
        segs = _image(src, ocr)
    elif ext in HTML_EXTS:
        segs = _html(src)
    elif ext in NOTEBOOK_EXTS:
        segs = _notebook(src)
    elif ext in TEXT_EXTS:
        segs = [("", _decode(_read_bytes(src)))]
    else:
        return []
    return [(loc, text) for loc, text in segs if text and text.strip()]


# ---------------------------------------------------------------- worker processes
_worker_ocr = False


def init_worker(ocr_enabled: bool, tesseract_cmd: str) -> None:
    """Process-pool initializer: configure OCR once per worker."""
    global _worker_ocr
    _worker_ocr = ocr_enabled and configure_ocr("auto", tesseract_cmd)


def extract_file(path: str) -> tuple[Segments, str]:
    """Process-pool entry point: returns (segments, error message)."""
    try:
        return extract(path, Path(path).suffix, ocr=_worker_ocr), ""
    except Exception as e:  # noqa: BLE001 - report any parser failure, keep going
        name = Path(path).name  # the full path is already shown next to the error
        msg = str(e).replace(repr(path)[1:-1], name).replace(path, name)
        return [], f"{type(e).__name__}: {msg}"
